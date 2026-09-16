import asyncio
import io
import json
import os
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[2]
PARSER_SERVICE = ROOT / "parser_service"
FIXTURE_XLSX = ROOT / "tests" / "fixtures" / "test.xlsx"
GOLDEN_PATH = Path(__file__).resolve().parent / "golden" / "ru_parser_100.json"

sys.path.insert(0, str(PARSER_SERVICE))

from parser_worker import process_excel_file

PARSE_DATE = datetime(2026, 9, 15)
LIMIT = 100
MOEX_LOSS_MAX = 2
INVESTING_LOSS_MAX = 5
ERROR_INCREASE_MAX = 3
MOEX_PRICE_TOL = 1.0
INVESTING_PRICE_TOL = 0.05
PRICE_REL_TOL = 0.02


def _num(value):
    if value is None or value == "":
        return None
    if value == "ERROR":
        return "ERROR"
    return float(value)


def snapshot_from_xlsx(content: bytes, date: datetime, limit: int) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    stocks = []
    for row in range(4, 4 + limit):
        if not ws.cell(row, 2).value:
            break
        stocks.append({
            "row": row,
            "isin": ws.cell(row, 2).value,
            "name": ws.cell(row, 3).value,
            "ticker": ws.cell(row, 15).value,
            "moex_price": _num(ws.cell(row, 7).value),
            "investing_price": _num(ws.cell(row, 8).value),
        })
    return {
        "date": date.strftime("%d.%m.%Y"),
        "limit": len(stocks),
        "stocks": stocks,
    }


def _counts(snapshot: dict) -> dict:
    stocks = snapshot["stocks"]
    moex = sum(1 for s in stocks if isinstance(s.get("moex_price"), float))
    investing = sum(1 for s in stocks if isinstance(s.get("investing_price"), float))
    errors = sum(1 for s in stocks if s.get("investing_price") == "ERROR")
    return {"moex": moex, "investing": investing, "errors": errors, "total": len(stocks)}


def _prices_close(expected, actual, abs_tol: float) -> bool:
    if expected is None or actual is None:
        return expected is None and actual is None
    if not isinstance(expected, float) or not isinstance(actual, float):
        return expected == actual
    if abs(expected - actual) <= abs_tol:
        return True
    denom = max(abs(expected), 1e-9)
    return abs(expected - actual) / denom <= PRICE_REL_TOL


def _compare(golden: dict, current: dict) -> list[str]:
    problems = []
    if golden["date"] != current["date"]:
        problems.append(f"date changed: golden {golden['date']}, got {current['date']}")
    if len(golden["stocks"]) != len(current["stocks"]):
        problems.append(
            f"row count changed: golden {len(golden['stocks'])}, got {len(current['stocks'])}"
        )
        return problems

    lost_moex = []
    lost_investing = []
    extra_errors = []
    price_drift = []

    for expected, actual in zip(golden["stocks"], current["stocks"]):
        label = actual.get("ticker") or actual.get("name")
        if expected.get("ticker") != actual.get("ticker"):
            problems.append(
                f"{label}: ticker changed {expected.get('ticker')} -> {actual.get('ticker')}"
            )
            continue

        if isinstance(expected.get("moex_price"), float) and not isinstance(actual.get("moex_price"), float):
            lost_moex.append(label)
        elif not _prices_close(expected.get("moex_price"), actual.get("moex_price"), MOEX_PRICE_TOL):
            price_drift.append(
                f"{label} MOEX {expected.get('moex_price')} -> {actual.get('moex_price')}"
            )

        if isinstance(expected.get("investing_price"), float) and not isinstance(actual.get("investing_price"), float):
            lost_investing.append(label)
        elif expected.get("investing_price") != "ERROR" and actual.get("investing_price") == "ERROR":
            extra_errors.append(label)
        elif isinstance(expected.get("investing_price"), float) and isinstance(actual.get("investing_price"), float):
            if not _prices_close(expected.get("investing_price"), actual.get("investing_price"), INVESTING_PRICE_TOL):
                price_drift.append(
                    f"{label} Investing {expected.get('investing_price')} -> {actual.get('investing_price')}"
                )

    golden_counts = _counts(golden)
    current_counts = _counts(current)

    if current_counts["moex"] < golden_counts["moex"] - MOEX_LOSS_MAX:
        problems.append(
            f"MOEX found {current_counts['moex']}/{current_counts['total']}, "
            f"golden {golden_counts['moex']} (max loss {MOEX_LOSS_MAX})"
        )
    if current_counts["investing"] < golden_counts["investing"] - INVESTING_LOSS_MAX:
        problems.append(
            f"Investing found {current_counts['investing']}/{current_counts['total']}, "
            f"golden {golden_counts['investing']} (max loss {INVESTING_LOSS_MAX})"
        )
    if current_counts["errors"] > golden_counts["errors"] + ERROR_INCREASE_MAX:
        problems.append(
            f"ERRORs {current_counts['errors']}, golden {golden_counts['errors']} "
            f"(max increase {ERROR_INCREASE_MAX})"
        )
    if golden_counts["investing"] > 0 and current_counts["investing"] == 0:
        problems.append("Investing.com found 0 prices while golden has hits")

    if len(lost_moex) > MOEX_LOSS_MAX:
        problems.append(f"lost MOEX prices ({len(lost_moex)}): {', '.join(lost_moex[:10])}")
    if len(lost_investing) > INVESTING_LOSS_MAX:
        problems.append(
            f"lost Investing prices ({len(lost_investing)}): {', '.join(lost_investing[:10])}"
        )
    if len(extra_errors) > ERROR_INCREASE_MAX:
        problems.append(f"new Investing ERRORs ({len(extra_errors)}): {', '.join(extra_errors[:10])}")
    if len(price_drift) > 3:
        problems.append(f"price drift: {'; '.join(price_drift[:8])}")

    return problems


@pytest.mark.integration
def test_parser_does_not_regress_on_sample_file():
    assert FIXTURE_XLSX.exists(), f"missing fixture: {FIXTURE_XLSX}"

    file_content = FIXTURE_XLSX.read_bytes()
    result_content, _summary = asyncio.run(
        process_excel_file(file_content, PARSE_DATE, limit=LIMIT)
    )
    current = snapshot_from_xlsx(result_content, PARSE_DATE, LIMIT)

    if os.getenv("UPDATE_GOLDEN") == "1":
        GOLDEN_PATH.parent.mkdir(parents=True, exist_ok=True)
        GOLDEN_PATH.write_text(
            json.dumps(current, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        pytest.skip(f"golden snapshot updated: {GOLDEN_PATH}")

    assert GOLDEN_PATH.exists(), (
        f"missing golden file {GOLDEN_PATH}. "
        "Run UPDATE_GOLDEN=1 pytest tests/integration/test_parser_regression.py"
    )
    golden = json.loads(GOLDEN_PATH.read_text(encoding="utf-8"))
    problems = _compare(golden, current)
    if problems:
        pytest.fail("parser regression vs golden:\n- " + "\n- ".join(problems))
